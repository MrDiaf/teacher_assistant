import asyncio
import io
import shutil
import uuid
from pathlib import Path

import pytest
from httpx import AsyncClient
from pptx import Presentation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.errors import CourseSourceFileTooLargeError
from app.course_sources.models import (
    CourseSource,
    CourseSourceType,
    ExtractedSourceText,
)
from app.course_sources.storage import CourseStorage
from tests.helpers import (
    create_calendar,
    create_course,
    create_event,
    upload_course_source,
)


def extractable_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


def presentation_bytes(text: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = text
    slide.placeholders[1].text = "Supporting slide text"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def stored_path(source: dict[str, object]) -> Path:
    settings = get_settings()
    return (
        settings.course_storage_root
        / str(source["course_id"])
        / str(source["source_type"])
        / str(source["id"])
        / str(source["stored_filename"])
    )


async def test_literature_and_slides_use_separate_storage_and_metadata(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    literature = await upload_course_source(
        client,
        course["id"],
        file_name="database-notes.txt",
        content=b"Normalization removes update anomalies.",
        source_type="literature",
        display_name="Database reading",
    )
    slides = await upload_course_source(
        client,
        course["id"],
        file_name="lecture-01.md",
        content=b"# Transactions\n\nAtomicity keeps changes together.",
        source_type="slides",
    )

    assert literature["source_type"] == "literature"
    assert slides["source_type"] == "slides"
    assert literature["display_name"] == "Database reading"
    assert slides["display_name"] == "lecture-01"
    assert literature["extraction_status"] == "completed"
    assert slides["extraction_status"] == "completed"
    assert literature["extracted_text_count"] == 1
    assert slides["extracted_text_count"] == 1
    assert "storage_path" not in literature
    assert stored_path(literature).read_bytes() == b"Normalization removes update anomalies."
    assert stored_path(slides).read_bytes().startswith(b"# Transactions")

    literature_list = await client.get(
        f"/api/v1/courses/{course['id']}/sources",
        params={"source_type": "literature"},
    )
    assert literature_list.status_code == 200
    assert [source["id"] for source in literature_list.json()] == [literature["id"]]

    extracted = list(
        await session.scalars(
            select(ExtractedSourceText).order_by(
                ExtractedSourceText.page_or_slide_number
            )
        )
    )
    assert {item.text for item in extracted} == {
        "Normalization removes update anomalies.",
        "# Transactions\n\nAtomicity keeps changes together.",
    }

    course_response = await client.get(f"/api/v1/courses/{course['id']}")
    assert course_response.json()["source_count"] == 2


async def test_pdf_and_pptx_extraction_preserves_page_and_slide_numbers(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    pdf = await upload_course_source(
        client,
        course["id"],
        file_name="reading.pdf",
        content=extractable_pdf("Relational algebra"),
        source_type="literature",
    )
    pptx = await upload_course_source(
        client,
        course["id"],
        file_name="lecture.pptx",
        content=presentation_bytes("Transport layer"),
        source_type="slides",
    )

    assert pdf["extraction_status"] == "completed"
    assert pdf["mime_type"] == "application/pdf"
    assert pptx["extraction_status"] == "completed"
    assert pptx["mime_type"].endswith("presentationml.presentation")

    rows = list(
        await session.scalars(
            select(ExtractedSourceText).order_by(ExtractedSourceText.source_id)
        )
    )
    assert all(row.page_or_slide_number == 1 for row in rows)
    assert any("Relational algebra" in row.text for row in rows)
    assert any("Transport layer" in row.text for row in rows)
    assert any("Supporting slide text" in row.text for row in rows)


async def test_unsupported_files_are_stored_and_spoofed_formats_are_rejected(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    unsupported = await upload_course_source(
        client,
        course["id"],
        file_name="diagram.bin",
        content=b"\x01\x02\x03course diagram",
        source_type="slides",
    )
    assert unsupported["extraction_status"] == "unsupported"
    assert unsupported["mime_type"] == "application/octet-stream"
    assert "unavailable" in unsupported["extraction_error"]
    assert stored_path(unsupported).is_file()

    invalid_pdf = await client.post(
        f"/api/v1/courses/{course['id']}/sources",
        params={
            "source_type": "literature",
            "file_name": "not-really.pdf",
        },
        content=b"This is not a PDF.",
    )
    traversal = await client.post(
        f"/api/v1/courses/{course['id']}/sources",
        params={
            "source_type": "literature",
            "file_name": "../escape.txt",
        },
        content=b"Blocked",
    )
    assert invalid_pdf.status_code == 400
    assert invalid_pdf.json()["code"] == "invalid_course_source_file"
    assert traversal.status_code == 400
    assert traversal.json()["code"] == "invalid_course_source_file"

    sources = list(await session.scalars(select(CourseSource)))
    assert [str(source.id) for source in sources] == [unsupported["id"]]


async def test_failed_extraction_can_retry_and_completed_extraction_is_idempotent(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    failed = await upload_course_source(
        client,
        course["id"],
        file_name="empty.txt",
        content=b"   \n  ",
    )
    completed = await upload_course_source(
        client,
        course["id"],
        file_name="useful.txt",
        content=b"Useful extracted text",
    )

    assert failed["extraction_status"] == "failed"
    assert "no readable text" in failed["extraction_error"]

    retry = await client.post(
        f"/api/v1/course-sources/{failed['id']}/retry-extraction"
    )
    assert retry.status_code == 200
    assert retry.json()["extraction_status"] == "failed"

    before = list(
        await session.scalars(
            select(ExtractedSourceText).where(
                ExtractedSourceText.source_id == uuid.UUID(completed["id"])
            )
        )
    )
    extract_again = await client.post(
        f"/api/v1/course-sources/{completed['id']}/extract"
    )
    assert extract_again.status_code == 200
    session.expire_all()
    after = list(
        await session.scalars(
            select(ExtractedSourceText).where(
                ExtractedSourceText.source_id == uuid.UUID(completed["id"])
            )
        )
    )
    assert [item.id for item in after] == [item.id for item in before]

    invalid_retry = await client.post(
        f"/api/v1/course-sources/{completed['id']}/retry-extraction"
    )
    assert invalid_retry.status_code == 409
    assert invalid_retry.json()["code"] == "course_source_extraction_state"


async def test_storage_enforces_streamed_size_limit(tmp_path: Path) -> None:
    storage = CourseStorage(tmp_path, maximum_bytes=5, maximum_mb=1)

    async def oversized_chunks():
        yield b"123"
        yield b"456"

    with pytest.raises(CourseSourceFileTooLargeError):
        await storage.store(
            course_id=uuid.uuid4(),
            source_id=uuid.uuid4(),
            source_type=CourseSourceType.LITERATURE,
            file_name="large.txt",
            chunks=oversized_chunks(),
            content_length=None,
        )
    assert not any(path.is_file() for path in tmp_path.rglob("*"))


async def test_source_deletion_is_isolated_and_missing_files_do_not_block_it(
    client: AsyncClient,
) -> None:
    course = await create_course(client)
    first = await upload_course_source(
        client,
        course["id"],
        file_name="first.txt",
        content=b"First source",
    )
    second = await upload_course_source(
        client,
        course["id"],
        file_name="second.txt",
        content=b"Second source",
    )

    deleted = await client.delete(f"/api/v1/course-sources/{first['id']}")
    assert deleted.status_code == 204
    assert not stored_path(first).exists()
    assert stored_path(second).is_file()

    await asyncio.to_thread(shutil.rmtree, stored_path(second).parent)
    missing_file_delete = await client.delete(
        f"/api/v1/course-sources/{second['id']}"
    )
    assert missing_file_delete.status_code == 204
    assert (
        await client.get(f"/api/v1/course-sources/{second['id']}")
    ).status_code == 404


async def test_course_deletion_removes_only_its_materials_and_preserves_calendar_data(
    client: AsyncClient,
) -> None:
    course = await create_course(client)
    other_course = await create_course(client, name="Other course")
    source = await upload_course_source(
        client,
        course["id"],
        file_name="private-reading.txt",
        content=b"Course-specific file",
    )
    other_source = await upload_course_source(
        client,
        other_course["id"],
        file_name="keep-reading.txt",
        content=b"Keep this file",
    )
    calendar = await create_calendar(client)
    event = await create_event(client, calendar["id"])

    deleted = await client.delete(f"/api/v1/courses/{course['id']}")
    assert deleted.status_code == 204
    assert not stored_path(source).exists()
    assert stored_path(other_source).is_file()
    assert (
        await client.get(f"/api/v1/calendars/{calendar['id']}")
    ).status_code == 200
    assert (await client.get(f"/api/v1/events/{event['id']}")).status_code == 200
