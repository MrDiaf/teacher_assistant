import io
import uuid

import pytest
from httpx import AsyncClient
from pptx import Presentation
from sqlalchemy.ext.asyncio import AsyncSession

from app.courses import public
from app.courses.public import (
    CourseContextQueryPort,
    CourseSummary,
    NoteContext,
    NoteExcerpt,
    NoteReference,
    SlideExcerpt,
    SlideReference,
    course_context_query_port,
)
from tests.helpers import (
    create_calendar,
    create_course,
    create_event,
    create_note,
    upload_course_source,
)


def deck_bytes(*titles: str) -> bytes:
    presentation = Presentation()
    for title in titles:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Supporting slide text"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_public_surface_is_read_only_and_exposes_no_persistence_types() -> None:
    assert set(public.__all__) == {
        "MAX_PUBLIC_CONTEXT_NOTES",
        "MAX_PUBLIC_CONTEXT_SLIDES",
        "MAX_PUBLIC_COURSES",
        "MAX_PUBLIC_COURSE_IDS",
        "MAX_PUBLIC_MATCH_TOKENS",
        "MAX_PUBLIC_NOTE_BODY_CHARS",
        "MAX_PUBLIC_NOTE_CONTEXT_CHARS",
        "MAX_PUBLIC_NOTE_REFERENCES",
        "MAX_PUBLIC_SLIDE_CONTEXT_CHARS",
        "MAX_PUBLIC_SLIDE_TEXT_CHARS",
        "CourseContextQueryPort",
        "CourseStatus",
        "CourseSummary",
        "NoteContext",
        "NoteExcerpt",
        "NoteReference",
        "SlideExcerpt",
        "SlideReference",
        "course_context_query_port",
    }
    operations = {
        name for name in vars(CourseContextQueryPort) if not name.startswith("_")
    }
    assert operations == {
        "find_notes",
        "get_note_context",
        "list_courses",
        "list_linked_calendar_ids",
    }
    for value in (
        CourseSummary,
        NoteReference,
        NoteExcerpt,
        SlideReference,
        SlideExcerpt,
        NoteContext,
    ):
        assert value.model_config["frozen"] is True
        assert not {"storage_path", "stored_filename", "content_hash"} & set(
            value.model_fields
        )
    # Discovery never returns a body; only the explicit excerpt query does.
    assert "body" not in NoteReference.model_fields


async def test_port_returns_typed_bounded_values_with_provenance(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    databases = await create_course(client, name="Databases", course_code="DAT101")
    networks = await create_course(client, name="Networks", course_code="NET200")
    slides = await upload_course_source(
        client,
        databases["id"],
        file_name="isolation.pptx",
        content=deck_bytes("Dirty reads", "Write skew", "Phantoms"),
        source_type="slides",
    )
    literature = await upload_course_source(
        client,
        databases["id"],
        file_name="reading.txt",
        content=b"Private literature about serializable isolation.",
        source_type="literature",
    )
    isolation = await create_note(
        client,
        databases["id"],
        title="Isolation levels",
        body="Serializable isolation prevents write skew.",
        slide_references=[
            {
                "source_id": slides["id"],
                "start_slide": 2,
                "end_slide": 2,
                "label": "Write skew slide",
            }
        ],
    )
    routing = await create_note(
        client, networks["id"], title="Routing", body="Dijkstra computes shortest paths."
    )

    port = course_context_query_port(session)

    courses = await port.list_courses()
    assert [course.name for course in courses] == ["Databases", "Networks"]
    assert all(isinstance(course, CourseSummary) for course in courses)
    assert courses[0].course_code == "DAT101"
    assert courses[0].status.value == "active"

    matched = await port.find_notes(match_tokens={"serializable"})
    assert [note.id for note in matched] == [uuid.UUID(isolation["id"])]
    assert matched[0].course_id == uuid.UUID(databases["id"])
    # Whole-word matching: a substring of a word is not a match.
    assert await port.find_notes(match_tokens={"serial"}) == ()
    assert await port.find_notes(match_tokens=set()) == ()
    scoped = await port.find_notes(course_ids=[uuid.UUID(networks["id"])])
    assert [note.id for note in scoped] == [uuid.UUID(routing["id"])]
    assert await port.find_notes(course_ids=[]) == ()

    context = await port.get_note_context([uuid.UUID(isolation["id"])])
    assert isinstance(context, NoteContext)
    assert [note.body for note in context.notes] == [
        "Serializable isolation prevents write skew."
    ]
    assert [
        (item.note_id, item.source_id, item.start_slide, item.end_slide, item.label)
        for item in context.slide_references
    ] == [
        (uuid.UUID(isolation["id"]), uuid.UUID(slides["id"]), 2, 2, "Write skew slide")
    ]
    # Only the referenced slide range is returned, with its source provenance.
    assert [(item.source_id, item.slide_number) for item in context.slide_text] == [
        (uuid.UUID(slides["id"]), 2)
    ]
    assert "Write skew" in context.slide_text[0].text
    # Literature is never exposed through the context contract.
    serialized = context.model_dump_json()
    assert literature["id"] not in serialized
    assert "Private literature" not in serialized
    assert "storage" not in serialized


async def test_port_enforces_character_and_count_budgets(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    first = await create_note(client, course["id"], title="First", body="a" * 50)
    second = await create_note(client, course["id"], title="Second", body="b" * 50)
    port = course_context_query_port(session)
    note_ids = [uuid.UUID(first["id"]), uuid.UUID(second["id"])]

    truncated = await port.get_note_context(note_ids, max_note_body_chars=20)
    assert [len(note.body) for note in truncated.notes] == [20, 20]

    budgeted = await port.get_note_context(
        note_ids, max_note_body_chars=40, max_note_context_chars=40
    )
    # The shared budget is exhausted by the first note, so the second is omitted.
    assert [(note.title, len(note.body)) for note in budgeted.notes] == [("First", 40)]

    assert len(await port.find_notes(course_ids=[uuid.UUID(course["id"])], limit=1)) == 1
    assert (await port.get_note_context([])).notes == ()
    unknown = await port.get_note_context([uuid.uuid4()])
    assert unknown == NoteContext(notes=(), slide_references=(), slide_text=())


@pytest.mark.parametrize(
    ("call", "message"),
    [
        (lambda port: port.list_courses(limit=0), "a course limit"),
        (
            lambda port: port.list_courses(limit=public.MAX_PUBLIC_COURSES + 1),
            "a course limit",
        ),
        (
            lambda port: port.find_notes(limit=public.MAX_PUBLIC_NOTE_REFERENCES + 1),
            "a note limit",
        ),
        (
            lambda port: port.find_notes(
                match_tokens={f"token{i}" for i in range(public.MAX_PUBLIC_MATCH_TOKENS + 1)}
            ),
            "match tokens",
        ),
        (
            lambda port: port.find_notes(
                course_ids=[uuid.uuid4() for _ in range(public.MAX_PUBLIC_COURSE_IDS + 1)]
            ),
            "course IDs",
        ),
        (
            lambda port: port.get_note_context(
                [uuid.uuid4() for _ in range(public.MAX_PUBLIC_CONTEXT_NOTES + 1)]
            ),
            "note IDs",
        ),
        (
            lambda port: port.get_note_context(
                [uuid.uuid4()],
                max_note_body_chars=public.MAX_PUBLIC_NOTE_BODY_CHARS + 1,
            ),
            "a note body size",
        ),
        (
            lambda port: port.get_note_context(
                [uuid.uuid4()], max_slides=public.MAX_PUBLIC_CONTEXT_SLIDES + 1
            ),
            "a slide limit",
        ),
        (
            lambda port: port.get_note_context(
                [uuid.uuid4()],
                max_slide_context_chars=public.MAX_PUBLIC_SLIDE_CONTEXT_CHARS + 1,
            ),
            "a slide context size",
        ),
        (
            lambda port: port.list_linked_calendar_ids(
                [uuid.uuid4() for _ in range(public.MAX_PUBLIC_COURSE_IDS + 1)]
            ),
            "course IDs",
        ),
    ],
)
async def test_port_rejects_requests_beyond_its_hard_limits(
    session: AsyncSession, call, message: str
) -> None:
    with pytest.raises(ValueError, match=message):
        await call(course_context_query_port(session))


async def test_linked_calendar_ids_are_associations_only_and_survive_course_deletion(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    other = await create_course(client, name="Other course", course_code="OTH1")
    lectures = await create_calendar(client, name="Lectures")
    labs = await create_calendar(client, name="Labs")
    event = await create_event(client, lectures["id"])
    for calendar in (lectures, labs):
        linked = await client.post(
            f"/api/v1/courses/{course['id']}/calendars",
            json={"calendar_id": calendar["id"]},
        )
        assert linked.status_code == 201, linked.text

    port = course_context_query_port(session)
    assert await port.list_linked_calendar_ids([uuid.UUID(course["id"])]) == (
        uuid.UUID(lectures["id"]),
        uuid.UUID(labs["id"]),
    )
    assert await port.list_linked_calendar_ids([uuid.UUID(other["id"])]) == ()
    assert await port.list_linked_calendar_ids([]) == ()

    deleted = await client.delete(f"/api/v1/courses/{course['id']}")
    assert deleted.status_code == 204, deleted.text
    session.expire_all()
    assert await port.list_linked_calendar_ids([uuid.UUID(course["id"])]) == ()
    # Removing the course removes only the association; Calendar data is untouched.
    assert (await client.get(f"/api/v1/calendars/{lectures['id']}")).status_code == 200
    assert (await client.get(f"/api/v1/calendars/{labs['id']}")).status_code == 200
    assert (await client.get(f"/api/v1/events/{event['id']}")).status_code == 200
