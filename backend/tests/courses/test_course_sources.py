import asyncio
import io
import shutil
import uuid
from pathlib import Path

import pytest
from httpx import AsyncClient
from pptx import Presentation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.course_sources.models import (
    CourseSource,
    CourseSourceType,
    ExtractedSourceText,
)
from app.course_sources.storage import CourseStorage
from app.courses.errors import CourseSourceFileTooLargeError
from tests.helpers import (
    create_calendar,
    create_course,
    create_event,
    upload_course_source,
)


def extractable_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


def presentation_bytes(text: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = text
    slide.placeholders[1].text = "Supporting slide text"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def stored_path(source: dict[str, object]) -> Path:
    settings = get_settings()
    return (
        settings.course_storage_root
        / str(source["course_id"])
        / str(source["source_type"])
        / str(source["id"])
        / str(source["stored_filename"])
    )


async def test_literature_and_slides_use_separate_storage_and_metadata(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    literature = await upload_course_source(
        client,
        course["id"],
        file_name="database-notes.txt",
        content=b"Normalization removes update anomalies.",
        source_type="literature",
        display_name="Database reading",
    )
    slides = await upload_course_source(
        client,
        course["id"],
        file_name="lecture-01.md",
        content=b"# Transactions\n\nAtomicity keeps changes together.",
        source_type="slides",
    )

    assert literature["source_type"] == "literature"
    assert slides["source_type"] == "slides"
    assert literature["display_name"] == "Database reading"
    assert slides["display_name"] == "lecture-01"
    assert literature["extraction_status"] == "completed"
    assert slides["extraction_status"] == "completed"
    assert literature["extracted_text_count"] == 1
    assert slides["extracted_text_count"] == 1
    assert "storage_path" not in literature
    assert stored_path(literature).read_bytes() == b"Normalization removes update anomalies."
    assert stored_path(slides).read_bytes().startswith(b"# Transactions")

    literature_list = await client.get(
        f"/api/v1/courses/{course['id']}/sources",
        params={"source_type": "literature"},
    )
    assert literature_list.status_code == 200
    assert [source["id"] for source in literature_list.json()] == [literature["id"]]

    extracted = list(
        await session.scalars(
            select(ExtractedSourceText).order_by(
                ExtractedSourceText.page_or_slide_number
            )
        )
    )
    assert {item.text for item in extracted} == {
        "Normalization removes update anomalies.",
        "# Transactions\n\nAtomicity keeps changes together.",
    }

    course_response = await client.get(f"/api/v1/courses/{course['id']}")
    assert course_response.json()["source_count"] == 2


async def test_pdf_and_pptx_extraction_preserves_page_and_slide_numbers(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    pdf = await upload_course_source(
        client,
        course["id"],
        file_name="reading.pdf",
        content=extractable_pdf("Relational algebra"),
        source_type="literature",
    )
    pptx = await upload_course_source(
        client,
        course["id"],
        file_name="lecture.pptx",
        content=presentation_bytes("Transport layer"),
        source_type="slides",
    )

    assert pdf["extraction_status"] == "completed"
    assert pdf["mime_type"] == "application/pdf"
    assert pptx["extraction_status"] == "completed"
    assert pptx["mime_type"].endswith("presentationml.presentation")

    rows = list(
        await session.scalars(
            select(ExtractedSourceText).order_by(ExtractedSourceText.source_id)
        )
    )
    assert all(row.page_or_slide_number == 1 for row in rows)
    assert any("Relational algebra" in row.text for row in rows)
    assert any("Transport layer" in row.text for row in rows)
    assert any("Supporting slide text" in row.text for row in rows)


async def test_pdf_content_streams_privately_with_head_and_byte_ranges(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    pdf_bytes = extractable_pdf("Private PDF byte-range fixture")
    source = await upload_course_source(
        client,
        course["id"],
        file_name="Course reading (week 1).pdf",
        content=pdf_bytes,
        source_type="literature",
    )
    content_url = f"/api/v1/course-sources/{source['id']}/content"

    full = await client.get(content_url)
    assert full.status_code == 200
    assert full.content == pdf_bytes
    assert full.headers["content-type"] == "application/pdf"
    assert full.headers["content-disposition"] == (
        "inline; filename*=utf-8''Course%20reading%20%28week%201%29.pdf"
    )
    assert full.headers["accept-ranges"] == "bytes"
    assert full.headers["cache-control"] == "private, no-store, max-age=0"
    assert full.headers["cross-origin-resource-policy"] == "same-origin"
    assert full.headers["referrer-policy"] == "no-referrer"
    assert full.headers["x-content-type-options"] == "nosniff"

    head = await client.head(content_url)
    assert head.status_code == 200
    assert head.content == b""
    assert head.headers["content-length"] == str(len(pdf_bytes))
    assert head.headers["accept-ranges"] == "bytes"

    ranged = await client.get(content_url, headers={"Range": "bytes=10-49"})
    assert ranged.status_code == 206
    assert ranged.content == pdf_bytes[10:50]
    assert ranged.headers["content-range"] == f"bytes 10-49/{len(pdf_bytes)}"
    assert ranged.headers["content-length"] == "40"
    assert ranged.headers["x-content-type-options"] == "nosniff"

    stored = await session.get(CourseSource, uuid.UUID(source["id"]))
    assert stored is not None
    assert stored.extraction_status.value == "completed"
    assert stored.extracted_text_count == 1
    assert "content_hash" not in source


async def test_pdf_content_rejects_non_pdf_unknown_and_invalid_ids(
    client: AsyncClient,
) -> None:
    course = await create_course(client)
    text_source = await upload_course_source(
        client,
        course["id"],
        file_name="reading.txt",
        content=b"This is private course text.",
    )

    non_pdf = await client.get(f"/api/v1/course-sources/{text_source['id']}/content")
    unknown = await client.get(f"/api/v1/course-sources/{uuid.uuid4()}/content")
    invalid = await client.get("/api/v1/course-sources/not-a-uuid/content")
    traversal = await client.get("/api/v1/course-sources/%2E%2E/content")

    assert non_pdf.status_code == 415
    assert non_pdf.json()["code"] == "course_source_not_pdf"
    assert unknown.status_code == 404
    assert unknown.json()["code"] == "course_source_not_found"
    assert invalid.status_code == 422
    assert traversal.status_code in {404, 422}
    for response in (non_pdf, unknown, invalid, traversal):
        assert "storage_path" not in response.text
        assert str(get_settings().course_storage_root) not in response.text


async def test_pdf_content_fails_closed_for_missing_changed_or_tampered_storage(
    client: AsyncClient,
    session: AsyncSession,
    tmp_path: Path,
) -> None:
    course = await create_course(client)

    missing = await upload_course_source(
        client,
        course["id"],
        file_name="missing.pdf",
        content=extractable_pdf("Missing file"),
    )
    stored_path(missing).unlink()
    missing_response = await client.get(f"/api/v1/course-sources/{missing['id']}/content")

    changed = await upload_course_source(
        client,
        course["id"],
        file_name="changed.pdf",
        content=extractable_pdf("Changed file"),
    )
    changed_path = stored_path(changed)
    changed_bytes = bytearray(changed_path.read_bytes())
    changed_bytes[-1] = ord("!")
    changed_path.write_bytes(changed_bytes)
    changed_response = await client.get(f"/api/v1/course-sources/{changed['id']}/content")

    tampered = await upload_course_source(
        client,
        course["id"],
        file_name="tampered.pdf",
        content=extractable_pdf("Tampered metadata"),
    )
    secret = tmp_path / "must-not-be-returned.txt"
    secret.write_text("arbitrary host content", encoding="utf-8")
    stored = await session.get(CourseSource, uuid.UUID(tampered["id"]))
    assert stored is not None
    stored.storage_path = "../../must-not-be-returned.txt"
    await session.commit()
    tampered_response = await client.get(f"/api/v1/course-sources/{tampered['id']}/content")

    for response in (missing_response, changed_response, tampered_response):
        assert response.status_code == 409
        assert response.json()["code"] == "course_source_content_unavailable"
        assert "arbitrary host content" not in response.text
        assert "storage_path" not in response.text
        assert str(get_settings().course_storage_root) not in response.text


async def test_unsupported_files_are_stored_and_spoofed_formats_are_rejected(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    unsupported = await upload_course_source(
        client,
        course["id"],
        file_name="diagram.bin",
        content=b"\x01\x02\x03course diagram",
        source_type="slides",
    )
    assert unsupported["extraction_status"] == "unsupported"
    assert unsupported["mime_type"] == "application/octet-stream"
    assert "unavailable" in unsupported["extraction_error"]
    assert stored_path(unsupported).is_file()

    invalid_pdf = await client.post(
        f"/api/v1/courses/{course['id']}/sources",
        params={
            "source_type": "literature",
            "file_name": "not-really.pdf",
        },
        content=b"This is not a PDF.",
    )
    traversal = await client.post(
        f"/api/v1/courses/{course['id']}/sources",
        params={
            "source_type": "literature",
            "file_name": "../escape.txt",
        },
        content=b"Blocked",
    )
    assert invalid_pdf.status_code == 400
    assert invalid_pdf.json()["code"] == "invalid_course_source_file"
    assert traversal.status_code == 400
    assert traversal.json()["code"] == "invalid_course_source_file"

    sources = list(await session.scalars(select(CourseSource)))
    assert [str(source.id) for source in sources] == [unsupported["id"]]


async def test_failed_extraction_can_retry_and_completed_extraction_is_idempotent(
    client: AsyncClient,
    session: AsyncSession,
) -> None:
    course = await create_course(client)
    failed = await upload_course_source(
        client,
        course["id"],
        file_name="empty.txt",
        content=b"   \n  ",
    )
    completed = await upload_course_source(
        client,
        course["id"],
        file_name="useful.txt",
        content=b"Useful extracted text",
    )

    assert failed["extraction_status"] == "failed"
    assert "no readable text" in failed["extraction_error"]

    retry = await client.post(
        f"/api/v1/course-sources/{failed['id']}/retry-extraction"
    )
    assert retry.status_code == 200
    assert retry.json()["extraction_status"] == "failed"

    before = list(
        await session.scalars(
            select(ExtractedSourceText).where(
                ExtractedSourceText.source_id == uuid.UUID(completed["id"])
            )
        )
    )
    extract_again = await client.post(
        f"/api/v1/course-sources/{completed['id']}/extract"
    )
    assert extract_again.status_code == 200
    session.expire_all()
    after = list(
        await session.scalars(
            select(ExtractedSourceText).where(
                ExtractedSourceText.source_id == uuid.UUID(completed["id"])
            )
        )
    )
    assert [item.id for item in after] == [item.id for item in before]

    invalid_retry = await client.post(
        f"/api/v1/course-sources/{completed['id']}/retry-extraction"
    )
    assert invalid_retry.status_code == 409
    assert invalid_retry.json()["code"] == "course_source_extraction_state"


async def test_storage_enforces_streamed_size_limit(tmp_path: Path) -> None:
    storage = CourseStorage(tmp_path, maximum_bytes=5, maximum_mb=1)

    async def oversized_chunks():
        yield b"123"
        yield b"456"

    with pytest.raises(CourseSourceFileTooLargeError):
        await storage.store(
            course_id=uuid.uuid4(),
            source_id=uuid.uuid4(),
            source_type=CourseSourceType.LITERATURE,
            file_name="large.txt",
            chunks=oversized_chunks(),
            content_length=None,
        )
    assert not any(path.is_file() for path in tmp_path.rglob("*"))


async def test_source_deletion_is_isolated_and_missing_files_do_not_block_it(
    client: AsyncClient,
) -> None:
    course = await create_course(client)
    first = await upload_course_source(
        client,
        course["id"],
        file_name="first.txt",
        content=b"First source",
    )
    second = await upload_course_source(
        client,
        course["id"],
        file_name="second.txt",
        content=b"Second source",
    )

    deleted = await client.delete(f"/api/v1/course-sources/{first['id']}")
    assert deleted.status_code == 204
    assert not stored_path(first).exists()
    assert stored_path(second).is_file()

    await asyncio.to_thread(shutil.rmtree, stored_path(second).parent)
    missing_file_delete = await client.delete(
        f"/api/v1/course-sources/{second['id']}"
    )
    assert missing_file_delete.status_code == 204
    assert (
        await client.get(f"/api/v1/course-sources/{second['id']}")
    ).status_code == 404


async def test_course_deletion_removes_only_its_materials_and_preserves_calendar_data(
    client: AsyncClient,
) -> None:
    course = await create_course(client)
    other_course = await create_course(client, name="Other course")
    source = await upload_course_source(
        client,
        course["id"],
        file_name="private-reading.txt",
        content=b"Course-specific file",
    )
    other_source = await upload_course_source(
        client,
        other_course["id"],
        file_name="keep-reading.txt",
        content=b"Keep this file",
    )
    calendar = await create_calendar(client)
    event = await create_event(client, calendar["id"])

    deleted = await client.delete(f"/api/v1/courses/{course['id']}")
    assert deleted.status_code == 204
    assert not stored_path(source).exists()
    assert stored_path(other_source).is_file()
    assert (
        await client.get(f"/api/v1/calendars/{calendar['id']}")
    ).status_code == 200
    assert (await client.get(f"/api/v1/events/{event['id']}")).status_code == 200
