from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

PDF_MIME_TYPE = "application/pdf"
PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
TEXT_MIME_TYPES = {"text/plain", "text/markdown"}


class ExtractionFailure(Exception):
    pass


@dataclass(frozen=True)
class ExtractedSegment:
    page_or_slide_number: int
    text: str
    content_hash: str


class TextExtractionAdapter:
    def supports(self, mime_type: str) -> bool:
        return mime_type in {PDF_MIME_TYPE, PPTX_MIME_TYPE, *TEXT_MIME_TYPES}

    def extract(self, path: Path, mime_type: str) -> list[ExtractedSegment]:
        try:
            if mime_type == PDF_MIME_TYPE:
                segments = self._extract_pdf(path)
            elif mime_type == PPTX_MIME_TYPE:
                segments = self._extract_pptx(path)
            elif mime_type in TEXT_MIME_TYPES:
                segments = self._extract_text(path)
            else:
                raise ExtractionFailure(
                    "Text extraction is unavailable for this file type."
                )
        except ExtractionFailure:
            raise
        except Exception as exc:
            raise ExtractionFailure(
                "Text could not be extracted from this file."
            ) from exc

        usable = [segment for segment in segments if segment.text.strip()]
        if not usable:
            raise ExtractionFailure(
                "The file contains no readable text. It may be image-only or empty."
            )
        return usable

    def _extract_pdf(self, path: Path) -> list[ExtractedSegment]:
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise ExtractionFailure(
                "Encrypted PDF files cannot be extracted. Upload an unlocked copy."
            )
        segments: list[ExtractedSegment] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                segments.append(self._segment(page_number, text))
        return segments

    def _extract_pptx(self, path: Path) -> list[ExtractedSegment]:
        presentation = Presentation(path)
        segments: list[ExtractedSegment] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            blocks: list[str] = []
            for shape in slide.shapes:
                blocks.extend(self._shape_text(shape))
            text = "\n".join(dict.fromkeys(block.strip() for block in blocks if block.strip()))
            if text:
                segments.append(self._segment(slide_number, text))
        return segments

    def _extract_text(self, path: Path) -> list[ExtractedSegment]:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ExtractionFailure(
                "Text and Markdown files must use UTF-8 encoding."
            ) from exc
        return [self._segment(1, text)] if text.strip() else []

    def _shape_text(self, shape: object) -> list[str]:
        blocks: list[str] = []
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            for child in getattr(shape, "shapes", []):
                blocks.extend(self._shape_text(child))
            return blocks
        if getattr(shape, "has_text_frame", False):
            text = getattr(shape, "text", "")
            if text:
                blocks.append(str(text))
        if getattr(shape, "has_table", False):
            table = shape.table  # type: ignore[attr-defined]
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        blocks.append(cell.text)
        return blocks

    @staticmethod
    def _segment(page_or_slide_number: int, text: str) -> ExtractedSegment:
        normalized = text.replace("\x00", "").strip()
        return ExtractedSegment(
            page_or_slide_number=page_or_slide_number,
            text=normalized,
            content_hash=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        )
